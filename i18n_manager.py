#!/usr/bin/env python3
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    print("Please install openpyxl: pip install openpyxl")
    sys.exit(1)

EXCEL_FILE = "ui_translations.xlsx"
MESSAGES_DIR = "frontend/messages"

def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Translations"
    ws.append(["Key", "en", "vi", "ja"])
    
    # Header
    ws.append(["header_badge", "AI Translation Engine", "Công Cụ Dịch AI", "AI翻訳エンジン"])
    ws.append(["header_title", "Document Translation", "Dịch Tài Liệu", "文書翻訳"])
    ws.append(["header_subtitle", "Multilingual document translation powered by AI", "Dịch tài liệu đa ngôn ngữ bằng AI", "AIを活用した多言語文書翻訳"])
    
    # Language Bar
    ws.append(["lang_source", "Source Language", "Ngôn Ngữ Nguồn", "翻訳元言語"])
    ws.append(["lang_target", "Target Language", "Ngôn Ngữ Đích", "翻訳先言語"])
    ws.append(["lang_auto_detect", "Auto Detect", "Tự Động Nhận Diện", "自動検出"])
    ws.append(["lang_domain", "Domain", "Chuyên Ngành", "ドメイン"])
    ws.append(["lang_swap_title", "Swap languages", "Hoán đổi ngôn ngữ", "言語を入れ替える"])
    
    # Upload Zone
    ws.append(["upload_title", "Drag & drop files or click to select", "Kéo thả file hoặc nhấn để chọn", "ファイルをドラッグ＆ドロップまたはクリックして選択"])
    ws.append(["upload_subtitle", "Supports common document formats", "Hỗ trợ các định dạng tài liệu phổ biến", "一般的な文書フォーマットをサポート"])
    ws.append(["upload_xliff_option", "Optional: Generate .xlf file for Review", "Tùy chọn: Sinh thêm file .xlf để Review", "オプション：レビュー用の.xlfファイルを生成"])
    ws.append(["upload_version", "Version", "Phiên bản", "バージョン"])
    ws.append(["upload_version_new", "Version: 2.1 (New)", "Phiên bản: 2.1 (Mới)", "バージョン：2.1（新）"])
    ws.append(["upload_version_old", "Version: 1.2 (Old - Trados/memoQ)", "Phiên bản: 1.2 (Cũ - Trados/memoQ)", "バージョン：1.2（旧 - Trados/memoQ）"])
    
    # Glossary
    ws.append(["glossary_title", "Glossary", "Thuật ngữ", "用語集"])
    ws.append(["glossary_replace_old", "Replace old", "Thay thế cũ", "古いものを置換"])
    ws.append(["glossary_upload_csv", "Upload CSV", "Tải CSV lên", "CSVをアップロード"])
    ws.append(["glossary_col_source", "Source", "Nguồn", "原文"])
    ws.append(["glossary_col_target", "Target", "Đích", "訳文"])
    ws.append(["glossary_col_context", "Context", "Ngữ cảnh", "コンテキスト"])
    ws.append(["glossary_col_action", "Action", "Thao tác", "操作"])
    ws.append(["glossary_empty", "No terms yet.", "Chưa có thuật ngữ nào.", "用語はまだありません。"])
    ws.append(["glossary_delete", "Delete", "Xóa", "削除"])
    ws.append(["glossary_delete_confirm", "Are you sure you want to delete this term?", "Bạn có chắc chắn muốn xóa thuật ngữ này?", "この用語を削除してもよろしいですか？"])
    ws.append(["glossary_uploading", "Uploading...", "Đang upload...", "アップロード中..."])
    
    # Job Cards
    ws.append(["job_queued", "Queued", "Đang chờ", "待機中"])
    ws.append(["job_extracting", "Extracting...", "Đang trích xuất...", "抽出中..."])
    ws.append(["job_translating", "Translating...", "Đang dịch...", "翻訳中..."])
    ws.append(["job_reconstructing", "Reconstructing...", "Đang tái tạo...", "再構築中..."])
    ws.append(["job_completed", "Completed", "Hoàn thành", "完了"])
    ws.append(["job_failed", "Failed", "Thất bại", "失敗"])
    ws.append(["job_download", "Download", "Tải về", "ダウンロード"])
    ws.append(["job_review", "Review", "Review", "レビュー"])
    ws.append(["job_download_xliff", "Download XLIFF", "Tải XLIFF", "XLIFFをダウンロード"])
    ws.append(["job_upload_xliff", "Upload translated XLIFF", "Upload XLIFF đã dịch", "翻訳済みXLIFFをアップロード"])
    ws.append(["job_details", "Details", "Chi tiết", "詳細"])
    
    # Editor
    ws.append(["editor_title", "Review & Edit Translation", "Review & Chỉnh Sửa Bản Dịch", "翻訳のレビューと編集"])
    ws.append(["editor_close", "Close Editor", "Đóng Editor", "エディターを閉じる"])
    ws.append(["editor_save", "Save & Export", "Lưu & Xuất File", "保存＆エクスポート"])
    ws.append(["editor_confidence", "Confidence", "Độ tự tin", "信頼度"])
    
    wb.save(EXCEL_FILE)
    print(f"Template populated and saved to {EXCEL_FILE}.")

    # Now generate JSON
    headers = [cell.value for cell in ws[1]]
    langs = [h for h in headers if h and h != "Key"]
    translations = {lang: {} for lang in langs}
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        key = row[0]
        for i, lang in enumerate(langs):
            val = row[i+1]
            if val is not None:
                translations[lang][key] = str(val)
                
    os.makedirs(MESSAGES_DIR, exist_ok=True)
    
    for lang, data in translations.items():
        out_path = os.path.join(MESSAGES_DIR, f"{lang}.json")
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        print(f"Generated {out_path} ({len(data)} keys)")

if __name__ == "__main__":
    main()
